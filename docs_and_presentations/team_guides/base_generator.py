import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "docs_and_presentations", "team_guides")
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Shared styling constants
PRIMARY_COLOR = RGBColor(15, 37, 55)       # Deep Navy #0F2537
SECONDARY_COLOR = RGBColor(24, 76, 120)    # Slate Blue #184C78
ACCENT_SKY = RGBColor(2, 132, 199)         # Sky Blue #0284C7
TEXT_DARK = RGBColor(31, 41, 55)           # Gray 800 #1F2937
TEXT_MUTED = RGBColor(75, 85, 99)          # Gray 600 #4B5563
EMERALD_GREEN = RGBColor(5, 150, 105)      # Emerald #059669
ROSE_RED = RGBColor(220, 38, 38)           # Red #DC2626

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTE / CRITICAL INSIGHT", border_color="0284C7", bg_color="F0F9FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.bold = True
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = SECONDARY_COLOR
    
    run_text = p.add_run(text)
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = TEXT_DARK
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(4)

def setup_document(title, subtitle, author, modules):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = TEXT_DARK
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Document Header
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("PRAKALP-DRISHTI")
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    run_sub = p_sub.add_run(f"{title}: {subtitle}")
    run_sub.font.size = Pt(15)
    run_sub.font.bold = True
    run_sub.font.color.rgb = ACCENT_SKY

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(12)
    run_meta = p_meta.add_run(
        f"Team Member: {author} | Assigned Modules: {modules}\n"
        f"Ministry of Statistics & Programme Implementation (MoSPI) | Smart India Hackathon 2026 (SIH26103)\n"
        f"Target: Complete Module Mastery, Mathematics, Data Provenance & Winning Jury Defense"
    )
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = TEXT_MUTED

    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(12)
    p_line_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="0284C7"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(p_line_border)

    return doc

def add_table(doc, headers, data):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "0F2537")
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(10)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)

    for row_idx, row_data in enumerate(data):
        row_cells = tbl.add_row().cells
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = str(text)
            set_cell_background(row_cells[col_idx], "F9FAFB" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.size = Pt(9.5)
            if col_idx == 0:
                p.runs[0].font.bold = True

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(4)

def add_qa_section(doc, qa_pairs):
    h = doc.add_heading("Winning Jury & Evaluator Combat Guide (Q&A)", level=1)
    h.style.font.color.rgb = PRIMARY_COLOR
    for q, a in qa_pairs:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        rq = p_q.add_run(f"❓ {q}")
        rq.bold = True
        rq.font.size = Pt(11)
        rq.font.color.rgb = PRIMARY_COLOR

        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_before = Pt(0)
        p_a.paragraph_format.space_after = Pt(6)
        ra = p_a.add_run(f"💬 {a}")
        ra.font.size = Pt(10)
        ra.font.color.rgb = TEXT_DARK

print("Shared helpers configured.")
